"""Export leads to CSV or Excel."""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("Name", "name"),
    ("Score", "score"),
    ("Qualified", "qualified"),
    ("Category", "category"),
    ("Best email", "best_email"),
    ("All emails", "all_emails"),
    ("Phone", "phone"),
    ("Website", "website"),
    ("Address", "address"),
    ("Rating", "rating"),
    ("Reviews", "reviews"),
    ("Matched keywords", "matched_keywords"),
    ("Why not qualified", "reasons"),
    ("Status", "status"),
    ("Notes", "notes"),
    ("Google Maps URL", "maps_url"),
    ("Search", "query"),
    ("Website title", "site_title"),
    ("Website problem", "site_error"),
    ("Latitude", "latitude"),
    ("Longitude", "longitude"),
]


def flatten(lead: dict) -> dict:
    emails = lead.get("emails") or []
    keywords = [
        f"{m['term']} ({m['count']})" for m in lead.get("matches") or [] if not m["term"].startswith("NOT: ")
    ]
    row = {key: lead.get(key) for _, key in COLUMNS}
    row.update(
        best_email=emails[0] if emails else "",
        all_emails=", ".join(emails),
        matched_keywords=", ".join(keywords),
        reasons="; ".join(lead.get("reasons") or []),
        qualified="yes" if lead.get("qualified") else "no",
    )
    return {k: ("" if v is None else v) for k, v in row.items()}


def to_csv(leads: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for label, _ in COLUMNS])
    for lead in leads:
        row = flatten(lead)
        writer.writerow([row[key] for _, key in COLUMNS])
    # BOM so Excel opens UTF-8 correctly.
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(leads: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([label for label, _ in COLUMNS])
    header_fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for lead in leads:
        row = flatten(lead)
        ws.append([row[key] for _, key in COLUMNS])
    widths = {"name": 32, "category": 22, "best_email": 30, "all_emails": 40, "website": 34, "address": 40,
              "matched_keywords": 40, "reasons": 34, "maps_url": 30, "notes": 30, "query": 28, "site_title": 30}
    for idx, (_, key) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(key, 12)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
