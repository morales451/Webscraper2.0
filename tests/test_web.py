import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from prospector import pipeline
from prospector.web.app import create_app

from .conftest import FakeSource, sample_places, site_handler


@pytest.fixture
def client(settings, monkeypatch):
    import httpx

    monkeypatch.setattr(
        pipeline, "make_http_client",
        lambda timeout=15.0: httpx.AsyncClient(transport=httpx.MockTransport(site_handler), follow_redirects=True),
    )
    monkeypatch.setattr(
        pipeline, "make_source",
        lambda *a, **k: FakeSource({"dentist in Austin, TX": sample_places()}),
    )
    with TestClient(create_app(settings)) as c:
        yield c


PROFILE = {
    "name": "Implant dentists",
    "queries": ["dentist"],
    "locations": ["Austin, TX"],
    "keywords": ["implant* | 3", "sedation | 2", "invisalign"],
    "filters": {"min_rating": 4.0},
}


def wait_for_run(client, run_id, timeout=20):
    deadline = time.time() + timeout
    while time.time() < deadline:
        run = client.get(f"/api/runs/{run_id}").json()
        if not run["active"] and run["status"] in ("done", "failed", "cancelled"):
            return run
        time.sleep(0.1)
    raise AssertionError("run did not finish")


def test_index_and_static(client):
    assert "Prospector" in client.get("/").text
    assert client.get("/static/app.js").status_code == 200


def test_profile_crud(client):
    created = client.post("/api/profiles", json=PROFILE).json()
    assert created["id"] == "implant-dentists"
    assert created["keywords"][0] == {"term": "implant*", "weight": 3.0}
    dup = client.post("/api/profiles", json=PROFILE).json()
    assert dup["id"] == "implant-dentists-2"
    updated = client.put("/api/profiles/implant-dentists", json={**PROFILE, "description": "hi"}).json()
    assert updated["description"] == "hi"
    assert {p["id"] for p in client.get("/api/profiles").json()} == {"implant-dentists", "implant-dentists-2"}
    assert client.delete("/api/profiles/implant-dentists-2").json() == {"ok": True}
    assert client.get("/api/profiles/implant-dentists-2").status_code == 404
    bad = client.post("/api/profiles", json={"name": "x", "queries": []})
    assert bad.status_code == 422 and "queries" in bad.json()["detail"]


def test_settings_mask_key(client):
    s = client.put("/api/settings", json={"places_api_key": "AIzaSyABCDEFGHIJKLMNOP", "site_concurrency": 4}).json()
    assert s["has_places_api_key"] and s["places_api_key"] == "AIza…MNOP" and s["site_concurrency"] == 4
    # Sending the masked value back must not overwrite the real key.
    s = client.put("/api/settings", json={"places_api_key": "AIza…MNOP"}).json()
    assert s["places_api_key"] == "AIza…MNOP"
    assert client.put("/api/settings", json={"source": "nope"}).status_code == 422


def test_run_leads_update_and_export(client):
    run = client.post("/api/runs", json={"profile": PROFILE, "source": "browser"}).json()
    run = wait_for_run(client, run["id"])
    assert run["status"] == "done" and run["processed"] == 5

    data = client.get(f"/api/leads?run_id={run['id']}&qualified=true").json()
    names = [l["name"] for l in data["leads"]]
    assert names[0] == "Smile Dental" and "Low Rated Dental" not in names
    assert client.get(f"/api/leads?run_id={run['id']}&has_email=true").json()["total"] >= 1
    assert client.get("/api/leads?q=plain").json()["leads"][0]["name"] == "Plain Dental"

    lead_id = data["leads"][0]["id"]
    upd = client.patch(f"/api/leads/{lead_id}", json={"status": "contacted", "notes": "called Tue"}).json()
    assert upd["status"] == "contacted" and upd["notes"] == "called Tue"
    assert client.patch(f"/api/leads/{lead_id}", json={"status": "bogus"}).status_code == 422
    assert client.get(f"/api/leads/{lead_id}").json()["notes"] == "called Tue"

    csv = client.get(f"/api/leads/export?format=csv&run_id={run['id']}&qualified=true")
    assert csv.headers["content-type"].startswith("text/csv")
    text = csv.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Name,Score,Qualified")
    assert "hello@smile.example" in text

    xlsx = client.get(f"/api/leads/export?format=xlsx&run_id={run['id']}")
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws.max_row == 6 and ws["A1"].value == "Name"

    assert client.delete(f"/api/runs/{run['id']}").json() == {"ok": True}
    assert client.get(f"/api/leads?run_id={run['id']}").json()["total"] == 0


def test_api_source_requires_key(client):
    r = client.post("/api/runs", json={"profile": PROFILE, "source": "api"})
    assert r.status_code == 422 and "API key" in r.json()["detail"]


def test_bulk_presets_selection_and_chains(client):
    run = client.post("/api/runs", json={"profile": PROFILE, "source": "browser"}).json()
    run = wait_for_run(client, run["id"])
    leads = client.get(f"/api/leads?run_id={run['id']}").json()["leads"]
    ids = [l["id"] for l in leads[:2]]

    r = client.post("/api/leads/bulk", json={"ids": ids, "status": "contacted"}).json()
    assert r == {"updated": 2}
    assert client.get("/api/leads?status=contacted").json()["total"] == 2
    assert client.post("/api/leads/bulk", json={"ids": ids, "status": "nope"}).status_code == 422

    sel = client.get(f"/api/leads/export?ids={ids[0]}&format=csv").content.decode("utf-8-sig").splitlines()
    assert len(sel) == 2
    hub = client.get(f"/api/leads/export?run_id={run['id']}&preset=hubspot")
    assert hub.content.decode("utf-8-sig").startswith("Company name,")
    assert 'leads-run' in hub.headers["content-disposition"] and "hubspot" in hub.headers["content-disposition"]
    assert client.get("/api/leads/export?preset=bogus").status_code == 422

    # Smile Dental and Low Rated Dental share smile.example: grouped they become one row.
    grouped = client.get(f"/api/leads?run_id={run['id']}&group_chains=true").json()
    assert grouped["total"] == run["processed"] - 1
    assert max(l["chain_size"] for l in grouped["leads"]) == 2


def test_profile_schedule_shows_next_run(client):
    client.post("/api/profiles", json={**PROFILE, "name": "Sched", "schedule": "weekly"})
    item = next(p for p in client.get("/api/profiles").json() if p["id"] == "sched")
    assert item["schedule"] == "weekly" and item["next_run_at"].startswith("1970")


def test_settings_new_fields_and_masking(client):
    s = client.put("/api/settings", json={"anthropic_api_key": "sk-ant-abcdefghijklmnop", "search_delay": 4,
                                          "maps_concurrency": 99, "ai_model": "claude-haiku-5-5"}).json()
    assert s["has_anthropic_api_key"] and s["anthropic_api_key"] == "sk-a…mnop"
    assert s["search_delay"] == 4 and s["maps_concurrency"] == 8 and s["ai_model"] == "claude-haiku-5-5"
