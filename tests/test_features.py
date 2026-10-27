import io
import json
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import httpx
import pytest
from openpyxl import load_workbook

from prospector import ai as ai_module
from prospector import pipeline
from prospector.export import split_address, to_csv, to_xlsx
from prospector.geo import Area, Geocoder, split_area
from prospector.icp import ICPProfile, ProfileStore
from prospector.pipeline import RunOptions, execute_run, related_categories
from prospector.scheduler import due_profiles
from prospector.sources.base import Place
from prospector.storage import Store
from prospector.suggest import SuggestionCollector, ignore_words_for

from .conftest import FakeSource, sample_places, site_handler


# ---------------------------------------------------------------- geo
def test_split_area_grid():
    city = Area(30.0, -98.0, 30.3, -97.6, "Austin")  # ~38 × 33 km
    cells = split_area(city, max_cells=9)
    assert len(cells) == 9
    assert cells[0].label == "Austin — area 1/9"
    assert cells[0].south == 30.0 and cells[-1].north == pytest.approx(30.3)
    town = Area(30.0, -98.0, 30.02, -97.98, "Tiny")  # ~2 km: never split below min_cell_km
    assert len(split_area(town, max_cells=9)) == 1
    assert 10 <= city.maps_zoom() <= 13


async def test_geocoder_caches(tmp_path):
    calls = []

    def handler(request):
        calls.append(str(request.url))
        assert "Prospector" in request.headers["user-agent"]
        return httpx.Response(200, json=[{"boundingbox": ["30.0", "30.3", "-98.0", "-97.6"], "display_name": "Austin, TX"}])

    cache = tmp_path / "geo.json"
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        geo = Geocoder(cache, client)
        first = await geo.geocode("Austin, TX")
        second = await geo.geocode("austin, tx ")
    assert first.area.north == 30.3 and second.display_name == "Austin, TX" and len(calls) == 1
    assert (await Geocoder(cache, client=None).geocode("Austin, TX")).area.west == -98.0  # from disk


# ---------------------------------------------------------------- suggestions
def test_suggestions_rank_phrases_common_on_best_leads():
    col = SuggestionCollector(known_terms=["implant*"], ignore_words=ignore_words_for(["dentist"], ["Boise, ID"]))
    for i in range(5):
        col.add(f"good{i}", ["Sedation dentistry for anxious patients. Dental implants. Serving Boise."])
    for i in range(5):
        col.add(f"meh{i}", ["Cleanings and exams. Serving Boise. Dental implants."])
    terms = [s["term"] for s in col.suggest([f"good{i}" for i in range(5)])]
    assert "sedation dentistry" in terms or "anxious patients" in terms
    assert not any("implant" in t or "boise" in t for t in terms)  # known keyword / location words
    assert col.suggest(["good0", "good1"]) == []  # too few best leads to say anything


def test_related_categories():
    cats = [["Dentist", "Cosmetic dentist", "Orthodontist"], ["Cosmetic dentist", "Dental clinic"],
            ["Orthodontist"], ["Cosmetic dentist", "Dental clinic"], ["Endodontist"]]
    # "cosmetic dentist" contains the query word, and a one-off category isn't a pattern.
    assert related_categories(cats, ["dentist"]) == ["orthodontist", "dental clinic"]
    assert related_categories(cats[:2], ["dentist"]) == []


# ---------------------------------------------------------------- export
LEAD = {
    "name": "Acme Roofing", "score": 80, "qualified": True, "category": "Roofing contractor",
    "emails": ["dead@acme.dead", "office@acmeroof.com"],
    "email_checks": {"dead@acme.dead": "invalid", "office@acmeroof.com": "ok"},
    "phone": "+1 509-555-0100", "site_phones": ["+15095550100", "+15095550199"],
    "website": "https://acmeroof.com/", "domain": "acmeroof.com",
    "address": "1 Main St, Spokane, WA 99201, United States",
    "matches": [{"term": "TPO / single-ply", "count": 3}, {"term": "NOT: residential only", "count": 1}],
    "signals": ["WordPress"], "socials": {"facebook": "https://facebook.com/acme"},
    "status": "contacted", "ai_note": "Commercial TPO specialist with property-manager clients.",
}


def test_split_address():
    assert split_address(LEAD["address"]) == {"street": "1 Main St", "city": "Spokane", "state": "WA", "postal_code": "99201"}


def test_export_presets():
    std = to_csv([LEAD]).decode("utf-8-sig").splitlines()
    assert std[0].startswith("Name,Score,Qualified,AI fit,AI note")
    assert "office@acmeroof.com" in std[1] and "+15095550199" in std[1]  # best deliverable email, other phone

    hub = to_csv([LEAD], "hubspot").decode("utf-8-sig").splitlines()
    assert hub[0].startswith("Company name,Company Domain Name,Website URL,Email")
    assert "ATTEMPTED_TO_CONTACT" in hub[1] and "Spokane,WA,99201" in hub[1]

    no_email = {**LEAD, "emails": ["dead@acme.dead"], "email_checks": {"dead@acme.dead": "invalid"}}
    out = to_csv([LEAD, no_email], "outreach").decode("utf-8-sig").splitlines()
    assert len(out) == 2 and out[1].startswith("office@acmeroof.com,Acme Roofing")  # undeliverable lead skipped
    assert "Commercial TPO specialist" in out[1]

    ws = load_workbook(io.BytesIO(to_xlsx([LEAD], "hubspot"))).active
    assert ws["A1"].value == "Company name" and ws["A2"].value == "Acme Roofing"


# ---------------------------------------------------------------- scheduler
def test_due_profiles(tmp_path):
    profiles = ProfileStore(tmp_path / "p")
    store = Store(tmp_path / "db.sqlite")
    profiles.save(ICPProfile(name="Daily", queries=["a"], schedule="daily"), "daily")
    profiles.save(ICPProfile(name="Weekly", queries=["a"], schedule="weekly"), "weekly")
    profiles.save(ICPProfile(name="Manual", queries=["a"]), "manual")
    assert sorted(s for s, _p in due_profiles(profiles, store)) == ["daily", "weekly"]  # never run yet
    store.mark_scheduled_run("daily", 1)
    store.mark_scheduled_run("weekly", 2)
    assert due_profiles(profiles, store) == []
    later = datetime.now(timezone.utc) + timedelta(days=2)
    assert [s for s, _p in due_profiles(profiles, store, now=later)] == ["daily"]


# ---------------------------------------------------------------- AI notes
async def test_fit_noter_parses_structured_output(monkeypatch):
    noter = ai_module.FitNoter("sk-test", "claude-opus-5-5")
    seen = {}

    async def fake_create(**kwargs):
        seen.update(kwargs)
        return SimpleNamespace(stop_reason="end_turn", content=[
            SimpleNamespace(type="thinking"), SimpleNamespace(type="text", text=json.dumps({"fit_score": 140, "note": " Great fit. "}))])

    monkeypatch.setattr(noter.client.beta.messages, "create", fake_create)
    profile = ICPProfile(name="Implants", queries=["dentist"], keywords=["implant*"], description="High-ticket")
    result = await noter.note(profile, Place(name="Smile Co", category="Dentist"), None, ["implant*"])
    assert result == {"fit_score": 100, "note": "Great fit."}
    assert seen["model"] == "claude-opus-5-5" and seen["fallbacks"] == "default"
    assert seen["output_config"]["format"]["type"] == "json_schema"
    assert "Smile Co" in seen["messages"][0]["content"] and "High-ticket" in seen["messages"][0]["content"]
    await noter.close()


async def test_fit_noter_refusal_returns_none(monkeypatch):
    noter = ai_module.FitNoter("sk-test")

    async def fake_create(**kwargs):
        return SimpleNamespace(stop_reason="refusal", content=[])

    monkeypatch.setattr(noter.client.beta.messages, "create", fake_create)
    assert await noter.note(ICPProfile(name="x", queries=["q"]), Place(name="A"), None, []) is None
    await noter.close()


# ---------------------------------------------------------------- pipeline features
@pytest.fixture
def fake_env(monkeypatch):
    monkeypatch.setattr(
        pipeline, "make_http_client",
        lambda timeout=15.0: httpx.AsyncClient(transport=httpx.MockTransport(site_handler), follow_redirects=True),
    )
    holder = {}

    def install(places_by_query):
        src = FakeSource(places_by_query)
        holder["src"] = src
        monkeypatch.setattr(pipeline, "make_source", lambda *a, **k: src)
        return src

    return install


class FakeGeocoder:
    async def geocode(self, location):
        from prospector.geo import Place as GeoPlace
        return GeoPlace(location, location, Area(30.0, -98.0, 30.3, -97.6, location))

    async def close(self):
        pass


async def test_split_areas_searches_each_cell(settings, fake_env):
    src = fake_env({"dentist": sample_places()})
    store = Store(settings.db_path)
    prof = ICPProfile(name="t", queries=["dentist"], locations=["Austin, TX"], split_areas=True, max_areas=4)
    run_id = store.create_run(prof.model_dump(), None, "fake", 1)
    run = await execute_run(settings, store, run_id, prof, RunOptions(enrich=False), geocoder=FakeGeocoder())
    assert run["searches_total"] == 4 and run["searches_done"] == 4
    assert all(area is not None for _q, area in src.calls) and len(src.calls) == 4
    assert run["found"] == 5  # the same businesses in every cell are merged
    leads, _ = store.query_leads(run_id=run_id)
    assert leads[0]["query"].startswith("dentist in Austin, TX — area")


async def test_expand_queries_adds_related_category_searches(settings, fake_env):
    base = [Place(name=f"Clinic {i}", place_id=f"c{i}", category="Dentist",
                  categories=["Dentist", "Orthodontist"]) for i in range(4)]
    extra = [Place(name="Brace Yourself", place_id="o1", category="Orthodontist")]
    src = fake_env({"dentist in Austin, TX": base, "orthodontist in Austin, TX": extra})
    store = Store(settings.db_path)
    prof = ICPProfile(name="t", queries=["dentist"], locations=["Austin, TX"], expand_queries=True)
    run_id = store.create_run(prof.model_dump(), None, "fake", 1)
    run = await execute_run(settings, store, run_id, prof, RunOptions(enrich=False))
    assert [q for q, _a in src.calls] == ["dentist in Austin, TX", "orthodontist in Austin, TX"]
    assert run["searches_total"] == 2 and run["found"] == 5


async def test_ai_notes_and_suggestions_are_saved(settings, fake_env, monkeypatch):
    fake_env({"dentist in Austin, TX": sample_places()})

    class FakeNoter:
        def __init__(self, *a, **k):
            self.disabled_reason = None

        async def note(self, profile, place, site, matched):
            return {"fit_score": 77, "note": f"{place.name} does implants."}

        async def close(self):
            pass

    monkeypatch.setattr(ai_module, "FitNoter", FakeNoter)
    settings.anthropic_api_key = "sk-test"
    store = Store(settings.db_path)
    prof = ICPProfile(name="t", queries=["dentist"], locations=["Austin, TX"], keywords=["implant*"], ai_notes=True)
    run_id = store.create_run(prof.model_dump(), None, "fake", 1)
    run = await execute_run(settings, store, run_id, prof)
    leads = {l["name"]: l for l in store.query_leads(run_id=run_id)[0]}
    assert leads["Smile Dental"]["ai_note"] == "Smile Dental does implants." and leads["Smile Dental"]["ai_score"] == 77
    assert leads["Smile Dental"]["socials"] == {"facebook": "https://facebook.com/smile"}
    assert isinstance(run["suggestions"], list)


async def test_ai_notes_without_key_are_skipped_with_a_log_line(settings, fake_env):
    fake_env({"dentist in Austin, TX": sample_places()[:1]})
    store = Store(settings.db_path)
    prof = ICPProfile(name="t", queries=["dentist"], locations=["Austin, TX"], ai_notes=True)
    run_id = store.create_run(prof.model_dump(), None, "fake", 1)
    run = await execute_run(settings, store, run_id, prof)
    assert run["status"] == "done" and any("no Anthropic API key" in l["msg"] for l in run["log"])


async def test_captcha_pauses_and_retries_once(settings, fake_env, monkeypatch):
    from prospector.sources.base import BlockedError

    monkeypatch.setattr(pipeline, "BLOCK_PAUSE_SECONDS", 0)
    attempts = []

    class Flaky(FakeSource):
        async def search(self, query, limit, area=None, want=None):
            attempts.append(query)
            if len(attempts) == 1:
                raise BlockedError("captcha")
            async for p in super().search(query, limit, area, want):
                yield p

    src = Flaky({"dentist in Austin, TX": sample_places()})
    monkeypatch.setattr(pipeline, "make_source", lambda *a, **k: src)
    store = Store(settings.db_path)
    prof = ICPProfile(name="t", queries=["dentist"], locations=["Austin, TX"])
    run_id = store.create_run(prof.model_dump(), None, "fake", 1)
    run = await execute_run(settings, store, run_id, prof, RunOptions(enrich=False))
    assert attempts == ["dentist in Austin, TX"] * 2 and run["status"] == "done" and run["found"] == 5
    assert any("CAPTCHA" in l["msg"] for l in run["log"])


def test_group_chains_and_bulk_update(tmp_path):
    store = Store(tmp_path / "db.sqlite")
    run_id = store.create_run({"name": "x"}, None, "fake", 1)
    for i, (domain, score) in enumerate([("chain.com", 50), ("chain.com", 80), ("chain.com", 60), ("solo.com", 40), ("", 30)]):
        store.save_lead(run_id, {"place_key": f"k{i}", "name": f"L{i}", "domain": domain, "score": score, "qualified": True})
    leads, total = store.query_leads(run_id=run_id, group_chains=True)
    assert total == 3 and [(l["name"], l["chain_size"]) for l in leads] == [("L1", 3), ("L3", 1), ("L4", 1)]
    assert store.update_leads([leads[0]["id"], leads[1]["id"]], "contacted") == 2
    assert store.query_leads(run_id=run_id, status="contacted")[1] == 2
    with pytest.raises(ValueError):
        store.update_leads([1], "bogus")
